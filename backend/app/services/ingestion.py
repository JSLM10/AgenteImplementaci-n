"""Extraccion, fragmentacion y vectorizacion de documentos.

Implementa los pasos 2 a 5 del flujo del enunciado: extraer el contenido,
dividirlo en fragmentos, generar los embeddings y almacenarlos en la base
vectorial. Los pasos 1 (recibir) y 6 (confirmar) viven en el router.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import uuid
from pathlib import Path

from langchain_core.documents import Document as LCDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import settings
from app.vectorstore.store import agregar_fragmentos

logger = logging.getLogger(__name__)

EXTENSIONES_SOPORTADAS = {".pdf", ".txt", ".md", ".markdown", ".csv", ".docx"}


class FormatoNoSoportadoError(ValueError):
    pass


class DocumentoVacioError(ValueError):
    """El archivo se leyo pero no contenia texto extraible."""


# ─────────────────────────────────────────────────────────────────────────────
# Paso 2 — extraer el contenido
# ─────────────────────────────────────────────────────────────────────────────
def _extraer_pdf(ruta: Path) -> list[tuple[str, int | None]]:
    from pypdf import PdfReader

    reader = PdfReader(str(ruta))
    paginas: list[tuple[str, int | None]] = []
    for numero, pagina in enumerate(reader.pages, start=1):
        texto = pagina.extract_text() or ""
        if texto.strip():
            paginas.append((texto, numero))
    return paginas


def _extraer_docx(ruta: Path) -> list[tuple[str, int | None]]:
    import docx

    documento = docx.Document(str(ruta))
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    # Las tablas suelen llevar los datos que mas se consultan; sin esto se
    # perderian por completo.
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    texto = "\n".join(partes)
    return [(texto, None)] if texto.strip() else []


def _extraer_csv(ruta: Path) -> list[tuple[str, int | None]]:
    # Se convierte cada fila a "columna: valor" porque un CSV crudo se
    # fragmenta fatal: los encabezados quedan en un trozo y los datos en otro.
    contenido = ruta.read_text(encoding="utf-8", errors="replace")
    lector = csv.DictReader(io.StringIO(contenido))
    filas = []
    for fila in lector:
        partes = [f"{k}: {v}" for k, v in fila.items() if k and v]
        if partes:
            filas.append("; ".join(partes))
    if not filas:
        return [(contenido, None)] if contenido.strip() else []
    return [("\n".join(filas), None)]


def _extraer_texto_plano(ruta: Path) -> list[tuple[str, int | None]]:
    texto = ruta.read_text(encoding="utf-8", errors="replace")
    return [(texto, None)] if texto.strip() else []


def extraer_contenido(ruta: Path) -> list[tuple[str, int | None]]:
    """Devuelve una lista de (texto, numero de pagina o None)."""
    extension = ruta.suffix.lower()

    if extension not in EXTENSIONES_SOPORTADAS:
        raise FormatoNoSoportadoError(
            f"El formato '{extension}' no está soportado. "
            f"Formatos aceptados: {', '.join(sorted(EXTENSIONES_SOPORTADAS))}."
        )

    if extension == ".pdf":
        partes = _extraer_pdf(ruta)
    elif extension == ".docx":
        partes = _extraer_docx(ruta)
    elif extension == ".csv":
        partes = _extraer_csv(ruta)
    else:
        partes = _extraer_texto_plano(ruta)

    if not partes:
        raise DocumentoVacioError(
            "No se pudo extraer texto del archivo. Si es un PDF escaneado, "
            "necesitaría OCR, que no está habilitado en esta versión."
        )
    return partes


# ─────────────────────────────────────────────────────────────────────────────
# Pasos 3 a 5 — fragmentar, vectorizar y almacenar
# ─────────────────────────────────────────────────────────────────────────────
def _crear_splitter() -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        # Se corta primero por parrafo y luego por frase, para que un fragmento
        # no parta una idea por la mitad.
        separators=["\n\n", "\n", ". ", " ", ""],
        add_start_index=True,
    )


_RE_ENCABEZADO = re.compile(r"^(#{1,6})\s+(.+?)\s*$", re.MULTILINE)


def _encabezados_por_posicion(texto: str) -> list[tuple[int, str]]:
    """Posicion de cada encabezado Markdown y su texto, en orden."""
    return [
        (m.start(), m.group(2).strip().strip("*").strip())
        for m in _RE_ENCABEZADO.finditer(texto)
    ]


def _encabezado_activo(encabezados: list[tuple[int, str]], posicion: int) -> str | None:
    """Ultimo encabezado que aparece antes de esa posicion del texto."""
    activo = None
    for inicio, titulo in encabezados:
        if inicio > posicion:
            break
        activo = titulo
    return activo


def procesar_documento(
    ruta: Path,
    document_id: uuid.UUID,
    nombre_archivo: str,
    owner_id: uuid.UUID | None,
    es_publico: bool,
) -> int:
    """Procesa un archivo entero y devuelve el numero de fragmentos creados.

    Cada fragmento se guarda precedido de una linea con el documento y la
    seccion a la que pertenece. No es decorativo: el fragmento se incrusta
    junto a esa linea, asi que un trozo suelto de una tabla de sanciones sigue
    sabiendo que habla de sanciones. Sin ese contexto, los fragmentos cortos
    de documentos largos se vuelven practicamente irrecuperables.
    """
    partes = extraer_contenido(ruta)
    splitter = _crear_splitter()
    titulo_documento = Path(nombre_archivo).stem.replace("-", " ").replace("_", " ")

    documentos: list[LCDocument] = []
    for texto, pagina in partes:
        encabezados = _encabezados_por_posicion(texto)

        for trozo in splitter.create_documents([texto]):
            fragmento = trozo.page_content
            if not fragmento.strip():
                continue

            seccion = _encabezado_activo(encabezados, trozo.metadata.get("start_index", 0))
            cabecera = titulo_documento if seccion is None else f"{titulo_documento} — {seccion}"
            if pagina is not None:
                cabecera += f" (página {pagina})"

            metadatos = {
                "document_id": str(document_id),
                "nombre_archivo": nombre_archivo,
                # El filtro de visibilidad de la busqueda depende de estos dos
                # campos; owner_id va como cadena para que JSONB lo compare bien.
                "owner_id": str(owner_id) if owner_id else "",
                "es_publico": es_publico,
            }
            if pagina is not None:
                metadatos["pagina"] = pagina
            if seccion:
                metadatos["seccion"] = seccion

            documentos.append(
                LCDocument(page_content=f"{cabecera}\n{fragmento}", metadata=metadatos)
            )

    if not documentos:
        raise DocumentoVacioError("El documento no produjo ningún fragmento utilizable.")

    total = agregar_fragmentos(documentos, document_id)
    logger.info("Documento %s procesado: %d fragmentos", nombre_archivo, total)
    return total
